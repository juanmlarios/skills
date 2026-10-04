"""From the skill directory: python3 scripts/test_build_presentation.py"""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from pptx import Presentation

from build_presentation import TEMPLATE, build, catalogue, load_template, validate_content
from template_support import NS

EXAMPLE = Path(__file__).resolve().parent.parent / "references/example-content.json"


class PresentationBuilderTest(unittest.TestCase):
    def test_all_layouts_notes_bullets_and_no_overwrite(self):
        original_hash = hashlib.sha256(TEMPLATE.read_bytes()).digest()
        samples = Presentation(TEMPLATE.with_name("Teal_Editorial_Starter.pptx"))
        content = {
            "title": "Builder test", "presenter": "Test presenter", "footer": "Test footer",
            "slides": [{
                "layout": sample.slide_layout.name,
                "fields": {ph.name: ph.text for ph in sample.placeholders},
                "notes": "Test speaking notes.",
                "sources": ["Test source: supplied presenter notes"], "duration_minutes": 1.5,
            } for sample in samples.slides],
        }
        # Exercise an intentionally empty field as well as multiline bullets.
        content["slides"][0]["fields"]["Subtitle"] = ""
        with TemporaryDirectory() as directory:
            output = Path(directory) / "all-layouts.pptx"
            build(content, output)
            deck = Presentation(output)
            self.assertEqual(len(deck.slides), 10)
            self.assertEqual(len(deck.slide_layouts), 10)
            self.assertEqual(deck.core_properties.title, content["title"])
            footer = next(shape for shape in deck.slide_masters[0].shapes if shape.name == "EDIT HERE: presentation footer")
            self.assertEqual(footer.text, content["footer"])
            for slide, expected in zip(deck.slides, content["slides"]):
                self.assertEqual(slide.slide_layout.name, expected["layout"])
                self.assertEqual({ph.name: ph.text for ph in slide.placeholders}, expected["fields"])
                self.assertIn("Test speaking notes.", slide.notes_slide.notes_text_frame.text)
                self.assertIn("Test source:", slide.notes_slide.notes_text_frame.text)
                self.assertIn("1.5 minutes", slide.notes_slide.notes_text_frame.text)
                for ph in slide.placeholders:
                    self.assertIsNone(ph._element.find("p:spPr/a:xfrm", NS))
                    for bullet in ph._element.findall(".//a:buChar", NS):
                        prop = bullet.getparent()
                        self.assertEqual(bullet.get("char"), "■")
                        self.assertEqual(prop.find("a:buSzPct", NS).get("val"), "70000")
                        self.assertEqual(prop.find("a:buFont", NS).get("typeface"), "Arial")
                        if ph.name == "Left panel points":
                            self.assertIsNotNone(prop.find("a:buClrTx", NS))
                        else:
                            self.assertEqual(prop.find("a:buClr/a:schemeClr", NS).get("val"), "accent1")
            saved = output.read_bytes()
            with self.assertRaises(FileExistsError):
                build(content, output)
            self.assertEqual(output.read_bytes(), saved)
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).digest(), original_hash)

    def test_manifest_errors_and_example(self):
        layouts = catalogue(load_template())
        example = json.loads(EXAMPLE.read_text(encoding="utf-8"))
        validate_content(example, layouts)
        for change in [
            lambda c: c["slides"][0].update(layout="Missing layout"),
            lambda c: c["slides"][0]["fields"].pop("Presenter"),
            lambda c: c["slides"][0]["fields"].update(Unexpected="Extra"),
            lambda c: c["slides"][0]["fields"].update(Presenter=["Wrong type"]),
            lambda c: c["slides"][0].update(notes=[]),
            lambda c: c["slides"][0].update(sources="Not a list"),
            lambda c: c["slides"][0].update(duration_minutes=-1),
            lambda c: c["slides"][0].update(duration_minutes=float("nan")),
            lambda c: c["slides"][0].update(duration_minutes=True),
            lambda c: c.update(slides=[]),
        ]:
            invalid = deepcopy(example)
            change(invalid)
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                validate_content(invalid, layouts)
        with TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                build(example, Path(directory) / "wrong-extension.potx")
            build(example, Path(directory) / "example.pptx")


if __name__ == "__main__":
    unittest.main()
