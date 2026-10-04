"""Build a talk from approved JSON content and the bundled Teal Editorial POTX."""
import argparse
from copy import deepcopy
from io import BytesIO
import json
import math
from pathlib import Path
import zipfile

from pptx import Presentation

from template_support import NS, replace_text

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets/Teal_Editorial_Template.potx"


def load_template():
    # python-pptx reads PPTX, not POTX. Change only the MIME type in memory;
    # keep the real template, layouts, theme, and relationships unchanged.
    buffer = BytesIO()
    with zipfile.ZipFile(TEMPLATE) as source, zipfile.ZipFile(buffer, "w") as copy:
        for name in source.namelist():
            value = source.read(name)
            if name == "[Content_Types].xml":
                value = value.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            copy.writestr(name, value)
    buffer.seek(0)
    return Presentation(buffer)


def catalogue(deck):
    return {layout.name: [ph.name for ph in layout.placeholders] for layout in deck.slide_layouts}


def validate_content(content, layouts):
    if not isinstance(content, dict) or set(content) != {"title", "presenter", "footer", "slides"}:
        raise ValueError("Content must contain exactly title, presenter, footer, and slides")
    for key in ["title", "presenter", "footer"]:
        if not isinstance(content[key], str) or (key == "title" and not content[key].strip()):
            raise ValueError(f"{key} must be a string; title cannot be empty")
    if not isinstance(content["slides"], list) or not content["slides"]:
        raise ValueError("slides must be a non-empty list")
    for number, slide in enumerate(content["slides"], 1):
        if not isinstance(slide, dict) or not {"layout", "fields"} <= set(slide) or set(slide) - {"layout", "fields", "notes", "sources", "duration_minutes"}:
            raise ValueError(f"Slide {number}: use layout, fields, and optional notes, sources, duration_minutes")
        name = slide["layout"]
        if not isinstance(name, str) or name not in layouts:
            raise ValueError(f"Slide {number}: unknown layout {name!r}; use --layouts")
        fields = slide["fields"]
        if not isinstance(fields, dict) or set(fields) != set(layouts[name]):
            raise ValueError(f"Slide {number}: fields must match {layouts[name]}")
        if not all(isinstance(value, str) for value in fields.values()):
            raise ValueError(f"Slide {number}: every field must be a plain string")
        if not isinstance(slide.get("notes", ""), str):
            raise ValueError(f"Slide {number}: notes must be a string")
        sources = slide.get("sources", [])
        if not isinstance(sources, list) or not all(isinstance(source, str) and source.strip() for source in sources):
            raise ValueError(f"Slide {number}: sources must be a list of non-empty reference strings")
        duration = slide.get("duration_minutes")
        if "duration_minutes" in slide and (type(duration) not in [int, float] or not math.isfinite(duration) or duration < 0):
            raise ValueError(f"Slide {number}: duration_minutes must be a finite, non-negative number")


def speaker_notes(slide):
    parts = [slide.get("notes", "")]
    if "duration_minutes" in slide:
        parts.append(f"Estimated delivery time: {slide['duration_minutes']:g} minutes")
    if slide.get("sources"):
        parts.append("Sources:\n" + "\n".join(slide["sources"]))
    return "\n\n".join(part for part in parts if part)


def build(content, output):
    output = Path(output).resolve()
    if output.suffix.lower() != ".pptx":
        raise ValueError("Output must be a .pptx file")
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}; choose a new version filename")
    deck = load_template()
    validate_content(content, catalogue(deck))
    # python-pptx has no public delete-slide API. Remove specimen slides and
    # their relationships; preserve the master and all ten reusable layouts.
    for identifier in list(deck.slides._sldIdLst):
        deck.part.drop_rel(identifier.rId)
        deck.slides._sldIdLst.remove(identifier)
    layouts = {layout.name: layout for layout in deck.slide_layouts}
    footer = next(shape for shape in deck.slide_masters[0].shapes if shape.name == "EDIT HERE: presentation footer")
    replace_text(footer._element, content["footer"])
    deck.core_properties.title = content["title"]
    deck.core_properties.author = content["presenter"]
    for item in content["slides"]:
        layout = layouts[item["layout"]]
        slide = deck.slides.add_slide(layout)
        originals = {ph.placeholder_format.idx: ph for ph in layout.placeholders}
        for ph in slide.placeholders:
            original = originals[ph.placeholder_format.idx]
            ph.name = original.name
            # Copy text formatting only. Geometry remains inherited, so native
            # Layout/Reset behaviour and master decoration are preserved.
            body = ph._element.find("p:txBody", NS)
            ph._element.replace(body, deepcopy(original._element.find("p:txBody", NS)))
            replace_text(ph._element, item["fields"][ph.name])
        notes = speaker_notes(item)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = BytesIO()
    deck.save(buffer)
    buffer.seek(0)
    checked = Presentation(buffer)
    if len(checked.slides) != len(content["slides"]) or len(checked.slide_layouts) != len(layouts):
        raise ValueError("Saved deck failed slide/layout count verification")
    for slide, item in zip(checked.slides, content["slides"]):
        if slide.slide_layout.name != item["layout"] or {ph.name: ph.text for ph in slide.placeholders} != item["fields"]:
            raise ValueError("Saved deck failed content/layout verification")
        notes = speaker_notes(item)
        if notes and slide.notes_slide.notes_text_frame.text != notes:
            raise ValueError("Saved deck failed speaker-note verification")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation prevents overwriting a concurrently created file.
    with output.open("xb") as target:
        try:
            target.write(buffer.getvalue())
        except BaseException:
            target.close()
            output.unlink()
            raise
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("content", type=Path, nargs="?", help="Approved content JSON")
    parser.add_argument("output", type=Path, nargs="?", help="New output PPTX; existing files are never overwritten")
    parser.add_argument("--layouts", action="store_true", help="List actual layout names and required text fields")
    args = parser.parse_args()
    try:
        if args.layouts:
            print(json.dumps(catalogue(load_template()), indent=2, ensure_ascii=False))
            return
        if args.content is None or args.output is None:
            parser.error("Supply content.json and output.pptx, or use --layouts")
        path = build(json.loads(args.content.read_text(encoding="utf-8")), args.output)
        print(f"Built and structurally checked {path}; render it before final delivery.")
    except (ValueError, OSError, zipfile.BadZipFile) as error:
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
